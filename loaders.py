"""Document loaders for PDF, PPTX, TXT, and MD files."""
import glob
import os
from io import BytesIO
from typing import List

from langchain_community.document_loaders import PyMuPDFLoader, TextLoader
from langchain_core.documents import Document
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import pytesseract


SUPPORTED_SUFFIXES = (".pdf", ".pptx", ".txt", ".md")


def load_pptx(path: str) -> List[Document]:
    """Extract text, tables, and OCR images from a PowerPoint file."""
    prs = Presentation(path)
    docs: List[Document] = []
    for i, slide in enumerate(prs.slides, 1):
        parts: List[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)

            if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                rows: List[str] = []
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    rows.append(" | ".join(cells))
                if rows:
                    parts.append("TABLE:\n" + "\n".join(rows))

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    img = Image.open(BytesIO(shape.image.blob))
                    ocr_text = pytesseract.image_to_string(img).strip()
                    if ocr_text:
                        parts.append(f"[IMAGE OCR]: {ocr_text}")
                except Exception:
                    pass

        content = "\n".join(parts)
        if content.strip():
            docs.append(Document(
                page_content=content,
                metadata={"source": path, "slide": i}
            ))
    return docs


def discover(docs_path: str) -> List[str]:
    """
    Every supported source file in docs_path, as absolute paths, sorted.

    Sorted so indexing order is deterministic between runs.
    """
    found: List[str] = []
    for suffix in SUPPORTED_SUFFIXES:
        found.extend(glob.glob(os.path.join(docs_path, f"*{suffix}")))
    return sorted(os.path.abspath(p) for p in found)


def load_file(path: str) -> List[Document]:
    """Load a single source file, dispatching on its extension."""
    suffix = os.path.splitext(path)[1].lower()
    if suffix == ".pdf":
        return PyMuPDFLoader(path).load()
    if suffix == ".pptx":
        return load_pptx(path)
    if suffix in (".txt", ".md"):
        return TextLoader(path, encoding="utf-8").load()
    raise ValueError(f"Unsupported file type: {path}")


def load_all(docs_path: str) -> List[Document]:
    """Load every supported document type from docs_path."""
    docs: List[Document] = []
    for path in discover(docs_path):
        docs.extend(load_file(path))
    return docs
